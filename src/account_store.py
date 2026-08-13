from __future__ import annotations

import csv
import uuid
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Mapping


@dataclass(frozen=True)
class AccountProfile:
    id: str
    name: str
    username: str
    password: str
    enabled: bool = True

    def to_dict(self) -> dict:
        return {
            'id': self.id,
            'name': self.name,
            'username': self.username,
            'password': self.password,
            'enabled': self.enabled,
        }


@dataclass(frozen=True)
class ImportedWorkspaceRow:
    account: AccountProfile
    task_title: str = ''
    task_url: str = ''
    # ``None`` keeps compatibility with callers that construct rows directly:
    # those callers get field inference, while parsed spreadsheet rows carry an
    # explicit (possibly empty) mask to preserve blank-cell semantics.
    provided_fields: frozenset[str] | None = None


_HEADER_ALIASES = {
    'name': {'账户名称', '账号名称', '名称', '备注', 'name', 'label'},
    'username': {'用户名', '账号', '账户', 'username', 'user'},
    'password': {'密码', 'password', 'pass', 'passwd'},
    'enabled': {'启用', '是否启用', 'enabled', 'active'},
    'task_title': {'任务名称', '课程名称', '标题', 'task', 'tasktitle'},
    'task_url': {'任务url', '课程url', 'url', 'taskurl', '链接'},
}


def _bool_value(value, default=True) -> bool:
    if value is None or str(value).strip() == '':
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() not in {
        '0', 'false', 'no', 'n', 'off', '否', '禁用', '停用',
    }


def _cell_is_provided(value) -> bool:
    """Return whether a spreadsheet cell contains an explicit update."""
    return value is not None and str(value).strip() != ''


def _text_value(value, *, strip: bool = False) -> str:
    text = '' if value is None else str(value)
    return text.strip() if strip else text


def normalize_accounts(config: Mapping) -> list[AccountProfile]:
    accounts = []
    seen_ids = set()
    for raw in config.get('accounts', []) or []:
        if not isinstance(raw, Mapping):
            continue
        username = str(raw.get('username', '')).strip()
        if not username:
            continue
        account_id = str(raw.get('id', '')).strip() or f'account-{uuid.uuid4().hex}'
        while account_id in seen_ids:
            account_id = f'account-{uuid.uuid4().hex}'
        seen_ids.add(account_id)
        accounts.append(AccountProfile(
            id=account_id,
            name=str(raw.get('name', '')).strip() or username,
            username=username,
            password=str(raw.get('password', '')),
            enabled=_bool_value(raw.get('enabled'), True),
        ))

    legacy_username = str(config.get('username', '')).strip()
    legacy_password = str(config.get('password', ''))
    if legacy_username.casefold() in {'用户名', 'username', 'user'} and (
        not legacy_password
        or legacy_password.casefold() in {'密码', 'password', 'pass'}
    ):
        legacy_username = ''
    if not accounts and legacy_username:
        accounts.append(AccountProfile(
            id='account-default',
            name=legacy_username,
            username=legacy_username,
            password=legacy_password,
            enabled=True,
        ))
    return accounts


def apply_account(config: Mapping, account: AccountProfile) -> dict:
    result = deepcopy(dict(config))
    result['username'] = account.username
    result['password'] = account.password
    result['account_id'] = account.id
    result['account_name'] = account.name
    return result


def merge_accounts(
    existing: Iterable[AccountProfile],
    incoming: Iterable[AccountProfile],
) -> list[AccountProfile]:
    result = list(existing)
    by_username = {item.username.casefold(): index for index, item in enumerate(result)}
    for item in incoming:
        key = item.username.casefold()
        if key not in by_username:
            by_username[key] = len(result)
            result.append(item)
            continue
        index = by_username[key]
        current = result[index]
        result[index] = AccountProfile(
            id=current.id,
            name=item.name or current.name,
            username=current.username,
            password=item.password or current.password,
            enabled=item.enabled,
        )
    return result


def merge_imported_accounts(
    existing: Iterable[AccountProfile],
    rows: Iterable[ImportedWorkspaceRow],
) -> list[AccountProfile]:
    """Merge spreadsheet rows without treating blank cells as updates."""
    result = list(existing)
    by_username = {item.username.casefold(): index for index, item in enumerate(result)}
    for row in rows:
        item = row.account
        key = item.username.casefold()
        if key not in by_username:
            by_username[key] = len(result)
            result.append(item)
            continue
        index = by_username[key]
        current = result[index]
        fields = row.provided_fields
        if fields is None:
            fields = frozenset(
                field for field, value in (
                    ('name', item.name),
                    ('password', item.password),
                )
                if _cell_is_provided(value)
            ) | {'enabled'}
        result[index] = AccountProfile(
            id=current.id,
            name=item.name if 'name' in fields else current.name,
            username=current.username,
            password=item.password if 'password' in fields else current.password,
            enabled=item.enabled if 'enabled' in fields else current.enabled,
        )
    return result


def _normalize_header(value) -> str:
    return ''.join(str(value or '').strip().lower().split()).replace('_', '')


def _canonical_headers(headers: Iterable) -> dict[str, str]:
    aliases = {
        canonical: {_normalize_header(item) for item in values}
        for canonical, values in _HEADER_ALIASES.items()
    }
    result = {}
    for header in headers:
        normalized = _normalize_header(header)
        for canonical, values in aliases.items():
            if normalized in values:
                result[str(header)] = canonical
                break
    return result


def _read_delimited(path: Path) -> list[dict]:
    content = None
    for encoding in ('utf-8-sig', 'utf-8', 'gb18030'):
        try:
            content = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    if content is None:
        raise ValueError('无法识别表格文本编码')
    delimiter = '\t' if path.suffix.lower() == '.tsv' else ','
    return list(csv.DictReader(content.splitlines(), delimiter=delimiter))


def _read_xlsx(path: Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError('读取 xlsx 需要安装 openpyxl') from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or '').strip() for value in next(rows, ())]
        return [
            {headers[index]: value for index, value in enumerate(row) if index < len(headers)}
            for row in rows
            if any(value not in (None, '') for value in row)
        ]
    finally:
        workbook.close()


def import_workspace_rows(path: str | Path) -> list[ImportedWorkspaceRow]:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix in {'.csv', '.tsv'}:
        raw_rows = _read_delimited(source)
    elif suffix == '.xlsx':
        raw_rows = _read_xlsx(source)
    else:
        raise ValueError('仅支持 .xlsx、.csv、.tsv 表格')
    if not raw_rows:
        return []

    header_map = _canonical_headers(raw_rows[0].keys())
    parsed_rows: list[tuple[str, str, str, frozenset[str]]] = []
    profiles: dict[str, AccountProfile] = {}
    for raw in raw_rows:
        values = {
            header_map[key]: value
            for key, value in raw.items()
            if key in header_map
        }
        username = _text_value(values.get('username'), strip=True)
        if not username:
            continue
        key = username.casefold()
        current = profiles.get(key)
        name = _text_value(values.get('name'), strip=True)
        password = _text_value(values.get('password'))
        enabled_value = values.get('enabled')
        provided_fields = frozenset(
            field for field in ('name', 'password', 'enabled')
            if field in values and _cell_is_provided(values[field])
        )
        profiles[key] = AccountProfile(
            id=current.id if current else f'account-{uuid.uuid4().hex}',
            name=name or (current.name if current else username),
            username=current.username if current else username,
            password=password or (current.password if current else ''),
            enabled=_bool_value(
                enabled_value,
                current.enabled if current else True,
            ),
        )
        parsed_rows.append((
            key,
            _text_value(values.get('task_title'), strip=True),
            _text_value(values.get('task_url'), strip=True),
            provided_fields,
        ))
    return [
        ImportedWorkspaceRow(
            account=profiles[key],
            task_title=task_title,
            task_url=task_url,
            provided_fields=provided_fields,
        )
        for key, task_title, task_url, provided_fields in parsed_rows
    ]


def export_import_template(path: str | Path) -> None:
    destination = Path(path)
    headers = ['账户名称', '用户名', '密码', '启用', '任务名称', '任务URL']
    if destination.suffix.lower() == '.csv':
        with destination.open('w', encoding='utf-8-sig', newline='') as file:
            writer = csv.writer(file)
            writer.writerow(headers)
            writer.writerow(['示例账号', 'student001', '', '是', '', ''])
        return
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise RuntimeError('导出 xlsx 需要安装 openpyxl') from exc
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '账号与任务'
    sheet.append(headers)
    sheet.append(['示例账号', 'student001', '', '是', '', ''])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='2563EB')
    widths = (18, 22, 22, 10, 32, 72)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = 'A2'
    workbook.save(destination)
